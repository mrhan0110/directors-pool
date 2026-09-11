"""XLSX 내보내기 테스트 (PRD F-02-3, F-09-17, 부록 C #9)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from data.models import Person
from data.session import session_scope
from reports.exports import candidate_rows, profile_statuses
from reports.xlsx import NOTICE, ExportMeta, build_table_xlsx

META = ExportMeta(
    title="후보 검색 결과",
    viewer_label="가상 담당자 (staff@example.com)",
    generated_at=datetime(2026, 9, 11, 10, 30),
    conditions=["성별: 여성", "연령: 55~59세"],
    total_matched=120,
    shown=30,
)


def _ids(n):
    with session_scope() as s:
        return list(s.execute(
            select(Person.person_id).where(Person.name_ko.like("가상%")).order_by(Person.person_id).limit(n)
        ).scalars())


def test_rows_keep_rank_order_and_have_source_urls():
    ids = list(reversed(_ids(5)))
    rows = candidate_rows(ids)
    assert [r["순위"] for r in rows] == [1, 2, 3, 4, 5]
    assert [r["성명"] for r in rows][0].startswith("가상")
    for r in rows:
        if r["현재 직업"] != "-" or r["타사 등기임원"] != "없음":
            assert "example.com" in r["출처 URL"]


def test_profile_statuses():
    ids = _ids(3)
    assert len(profile_statuses(ids)) == 3


def test_xlsx_requires_source_url_column():
    with pytest.raises(ValueError):
        build_table_xlsx([{"성명": "가상"}], META)


def test_xlsx_contains_meta_and_watermark():
    rows = candidate_rows(_ids(3))
    data = build_table_xlsx(rows, META)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["후보 목록", "조회 정보"]
    ws = wb["후보 목록"]
    assert "staff@example.com" in ws.oddHeader.center.text
    assert NOTICE in ws.oddFooter.center.text
    headers = [c.value for c in ws[4]]
    assert "출처 URL" in headers
    assert ws.cell(row=5, column=headers.index("성명") + 1).value == rows[0]["성명"]
    info = {r[0].value: r[1].value for r in wb["조회 정보"].iter_rows(min_row=2)}
    assert info["전체 매칭 인원"] == 120
    assert info["조회 인원 수(상위 N명)"] == 30
    assert "성별: 여성" in info["조회 조건"]
    assert info["조회 일시"].startswith("2026-09-11")
