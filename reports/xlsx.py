"""XLSX 내보내기 (PRD F-02-3, F-07, F-09-17, 부록 C #9).

- 목록 시트: 순위 순서 그대로 + 출처 URL 컬럼
- 조회 정보 시트: 조회 조건 · 전체 매칭 수 · 조회 인원 수 · 조회 일시 · 열람자
- 모든 시트 머리글/바닥글에 열람자·일시 워터마크와 대외비 고지
출력 가능 여부(검수 게이트·권한)는 호출 전에 reports.builder.check_export 로 판정한다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

NOTICE = "대외비 · 공개정보 기반 참고자료 · 최종 판단은 담당자·법무 검토로 확정"


@dataclass
class ExportMeta:
    title: str
    viewer_label: str                      # 예: "가상 담당자 (staff@example.com)"
    generated_at: datetime
    conditions: list[str] = field(default_factory=list)
    total_matched: int | None = None
    shown: int | None = None


def _watermark(ws, meta: ExportMeta) -> None:
    stamp = f"열람자 {meta.viewer_label} · {meta.generated_at:%Y-%m-%d %H:%M}"
    ws.oddHeader.center.text = f"대외비 — {stamp}"
    ws.oddFooter.center.text = NOTICE
    ws.evenHeader.center.text = f"대외비 — {stamp}"
    ws.evenFooter.center.text = NOTICE


def build_table_xlsx(rows: list[dict], meta: ExportMeta, sheet_title: str = "후보 목록") -> bytes:
    """rows 는 표시 순서대로 정렬된 dict 목록. '출처 URL' 키가 반드시 있어야 한다."""
    if rows and any("출처 URL" not in r for r in rows):
        raise ValueError("출처 URL 컬럼 없이 내보낼 수 없습니다. (PRD F-02-3)")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([f"{meta.title} — {NOTICE}"])
    ws["A1"].font = Font(bold=True)
    ws.append([f"열람자 {meta.viewer_label} · 생성 {meta.generated_at:%Y-%m-%d %H:%M}"])
    ws.append([])
    headers = list(rows[0].keys()) if rows else ["출처 URL"]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    for idx, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r.get(h) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(idx)].width = min(60, max(8, width + 2))
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"
    _watermark(ws, meta)

    info = wb.create_sheet("조회 정보")
    info.append(["항목", "내용"])
    info.append(["자료명", meta.title])
    info.append(["조회 조건", " | ".join(meta.conditions) if meta.conditions else "전체"])
    if meta.total_matched is not None:
        info.append(["전체 매칭 인원", meta.total_matched])
    if meta.shown is not None:
        info.append(["조회 인원 수(상위 N명)", meta.shown])
    info.append(["조회 일시", f"{meta.generated_at:%Y-%m-%d %H:%M:%S}"])
    info.append(["열람자", meta.viewer_label])
    info.append(["고지", NOTICE])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 90
    _watermark(info, meta)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
